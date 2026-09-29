#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extract ALL paragraphs with MATH markers from Word doc."""
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

doc = Document('Lossless Transmission for Geo-Distributed AI Computing via LSQ-Sketch and APN-Aware SRv6 Backpressure.docx')

print("=== ALL paragraphs containing MATH elements ===")

math_paras = []
for i, para in enumerate(doc.paragraphs):
    has_math = False
    for elem in para._element.iter():
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag in ('oMathPara', 'oMath'):
            has_math = True
            break
    if has_math:
        math_paras.append(i)

print(f"Math paragraphs found: {math_paras}")
print(f"Total {len(math_paras)} paragraphs\n")

for idx in math_paras:
    para = doc.paragraphs[idx]
    style = para.style.name if para.style else 'Normal'
    
    # Extract runs sequence
    runs_seq = []
    for elem in para._element:
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'r':
            run_text = ''.join(t.text or '' for t in elem.iter(qn('w:t')))
            if run_text.strip():
                runs_seq.append(('T', run_text))
        elif tag in ('oMathPara', 'oMath'):
            math_texts = [t.text for t in elem.iter(qn('m:t')) if t.text]
            struct = []
            for child in elem.iter():
                ctag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                if ctag in ('f', 'rad', 'sSup', 'sSub', 'sSubSup', 'd', 'nary', 'eqArr', 'bar'):
                    struct.append(ctag)
            runs_seq.append(('M', {'t': math_texts, 's': struct}))
    
    print(f"\n[P{idx:3d}] [{style}]")
    for item in runs_seq:
        if item[0] == 'T':
            # Text
            print(f"  TXT: {repr(item[1])}")
        else:
            m = item[1]
            latex_hint = ''
            # Try to guess latex based on structure
            if 'f' in m['s']:
                latex_hint = ' [FRACTION]'
            if 'sSub' in m['s']:
                latex_hint += ' [SUBSCRIPT]'
            if 'sSup' in m['s']:
                latex_hint += ' [SUPERSCRIPT]'
            if 'sSubSup' in m['s']:
                latex_hint += ' [SUB+SUP]'
            if 'rad' in m['s']:
                latex_hint += ' [SQRT]'
            if 'nary' in m['s']:
                latex_hint += ' [OPERATOR sum/product etc]'
            if 'd' in m['s']:
                latex_hint += ' [DELIMITERS parens etc]'
            print(f"  MATH: {m['t']}{latex_hint}")
