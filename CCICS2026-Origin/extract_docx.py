from docx import Document
import json

doc = Document('Lossless Transmission for Geo-Distributed AI Computing via LSQ-Sketch and APN-Aware SRv6 Backpressure.docx')

content = []
for i, para in enumerate(doc.paragraphs):
    style_name = para.style.name if para.style else 'Normal'
    text = para.text.strip()
    if text:
        runs_info = []
        for run in para.runs:
            if run.text.strip():
                runs_info.append({
                    'text': run.text,
                    'bold': run.bold,
                    'italic': run.italic,
                    'underline': run.underline
                })
        content.append({
            'index': i,
            'style': style_name,
            'text': text,
            'runs': runs_info
        })

tables = []
for t_idx, table in enumerate(doc.tables):
    table_data = []
    for row in table.rows:
        row_data = [cell.text.strip() for cell in row.cells]
        table_data.append(row_data)
    tables.append({
        'index': t_idx,
        'data': table_data
    })

result = {
    'paragraphs': content,
    'tables': tables
}

with open('extracted_content.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print(f"Extracted {len(content)} paragraphs and {len(tables)} tables")
